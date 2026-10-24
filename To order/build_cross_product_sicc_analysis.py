from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import math
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from matplotlib.ticker import LogFormatter, LogLocator, NullFormatter


DOMAIN_ALIAS_MAP = {
    "AT": "Atom",
    "ATA": "Atom",
    "ATOM": "Atom",
    "ATOM_IN": "Atom",
    "ATOM_OUT": "Atom",
    "ATSRAMA": "Atom",
    "ATSRAM": "Atom",
    "LPATOM": "Atom",
    "LPATOMB": "Atom",
    "LPATSRAM": "Atom",
    "CORE": "Core",
    "CORE_IN": "Core",
    "CORE_OUT": "Core",
    "CR": "Core",
    "IA": "Core",
    "IAA": "Core",
    "CCF": "Ring",
    "CCFA": "Ring",
    "CCFSRAMA": "Ring",
    "RING": "Ring",
    "GT": "GT",
    "GT_IN": "GT",
    "GT_OUT": "GT",
    "SAAT": "SA-Atom",
    "SAREST": "SA-Atom",
    "SAHUB": "SAHUB",
    "NPU": "NPU",
    "NPU_IN": "NPU",
    "NPU_OUT": "NPU",
    "NPUSRAM": "NPUSRAM",
    "IPU": "IPU",
    "MEDIA": "MEDIA",
    "MEDIA_IN": "MEDIA",
    "MEDIA_OUT": "MEDIA",
    "MEDIASRAM": "MEDIASRAM",
    "DE": "DE",
    "DESRAM": "DESRAM",
}

RAW_ALIAS_PRIORITY = {
    "Atom": ["ATOM_OUT"],
    "SA-Atom": ["ATOM_OUT"],
    "Core": ["CORE_OUT"],
    "Ring": ["CORE_OUT"],
    "GT": ["GT_OUT"],
    "NPU": ["NPU_OUT"],
    "NPUSRAM": ["NPU_OUT"],
    "MEDIA": ["MEDIA_OUT"],
    "MEDIASRAM": ["MEDIA_OUT"],
}

PRODUCT_COLORS = {
    "NVLAX": "#1f77b4",
    "NVLHX": "#ff7f0e",
    "NVLS7": "#2ca02c",
}

PRODUCT_DISPLAY_NAMES = {
    "NVLAX": "NVLAX",
    "NVLS7": "NVLS",
    "NVLHX": "NVLHX",
}

MEASURE_DISPLAY_NAMES = {
    "raw": "raw",
    "tester_at_ref": "tester at ref",
    "system_at_ref": "system at ref",
}

RAW_DOMAIN_LABELS = {
    "IA": "Core",
    "AT": "Atom",
    "CCF": "Ring",
    "GT": "GT",
    "SAAT": "SA-Atom",
    "SAHUB": "SAHUB",
    "SANPU": "NPU",
    "SA": "SA",
}


@dataclass(frozen=True)
class MeasureColumns:
    raw: str | None
    system_at_ref: str | None
    tester_at_ref: str | None


@dataclass(frozen=True)
class MetricConfig:
    canonical_domain: str
    domain_token: str
    measure_type: str
    source_column: str
    upm_column: str
    raw_mapping_note: str | None
    core_number: int | None


@dataclass(frozen=True)
class TokenRule:
    source_token: str
    canonical_domain: str
    core_number: int | None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build cross-product SICC analysis artifacts and PowerPoint output.")
    parser.add_argument(
        "--input-csv",
        default=r"R:\Products\NVL\NVL-AX\Analysis\2026_43_Cross product_SICC_analysis.csv",
        help="Cross-product source CSV path.",
    )
    parser.add_argument(
        "--presi-target-xlsx",
        default=r"R:\Products\NVL\NVL-AX\Weekly data pull\PreSi SICC targets.xlsx",
        help="Workbook holding PreSi SICC targets.",
    )
    parser.add_argument(
        "--token-dictionary-xlsx",
        default=r"R:\Products\NVL\NVL-AX\Weekly data pull\SICC token naming.xlsx",
        help="Workbook holding the authoritative SICC token include/mapping dictionary.",
    )
    parser.add_argument(
        "--output-dir",
        default=r"R:\Products\NVL\NVL-AX\Weekly data pull\2026_43_Cross_product_SICC_analysis",
        help="Directory for generated analysis artifacts.",
    )
    parser.add_argument(
        "--deck-path",
        default=r"R:\Products\NVL\NVL-AX\Weekly data pull\2026_43_Cross_product_SICC_analysis.pptx",
        help="Target PowerPoint path.",
    )
    parser.add_argument(
        "--profile-only",
        action="store_true",
        help="Only emit schema, target, and dedupe artifacts. Skip normalization, plots, and PPT output.",
    )
    parser.add_argument(
        "--chunk-size",
        type=int,
        default=25000,
        help="CSV chunk size for streamed loading.",
    )
    parser.add_argument(
        "--core-shift-threshold-pct",
        type=float,
        default=5.0,
        help="Relative spread threshold for flagging token/core shifts.",
    )
    parser.add_argument(
        "--program-shift-threshold-pct",
        type=float,
        default=5.0,
        help="Relative spread threshold for flagging per-program shifts within one product.",
    )
    return parser.parse_args()


def safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", value.strip())
    cleaned = re.sub(r"(^-+|-+$)", "", cleaned)
    return cleaned.lower() or "unknown"


def read_csv_header(input_csv: Path) -> list[str]:
    with input_csv.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        return next(reader)


def resolve_required_columns(columns: list[str]) -> dict[str, str]:
    upper_map = {column.upper(): column for column in columns}
    required = {
        "prod": "PROD",
        "visual_id": "VISUAL_ID",
        "program_name": "PROGRAM NAME_CLASSHOT",
        "end_time": "LOTS END DATE TIME_CLASSHOT",
    }
    resolved: dict[str, str] = {}
    for logical_name, upper_name in required.items():
        if upper_name not in upper_map:
            raise ValueError(f"Required column '{upper_name}' is missing from the source CSV.")
        resolved[logical_name] = upper_map[upper_name]
    return resolved


def extract_domain_token(column_name: str) -> str | None:
    summary_match = re.search(r"SUMMARY-RESULTS\[(?P<domain>[A-Z0-9_]+)(?:::.*)?\]_CLASSHOT$", column_name, re.IGNORECASE)
    if summary_match:
        return summary_match.group("domain")

    direct_match = re.search(r"VA-NA-(?P<domain>[A-Z0-9_]+)-SICC_(?:SYSTEM|TESTER)_AT_REF_CLASSHOT$", column_name, re.IGNORECASE)
    if direct_match:
        return direct_match.group("domain")

    rec_match = re.search(r"RECLOGIN-(?P<domain>[A-Z0-9_]+)_CLASSHOT$", column_name, re.IGNORECASE)
    if rec_match:
        return rec_match.group("domain")

    return None


def canonical_domain_name(domain_token: str) -> str:
    upper = domain_token.strip().upper()
    if upper in DOMAIN_ALIAS_MAP:
        return DOMAIN_ALIAS_MAP[upper]
    stripped = re.sub(r"\d+$", "", upper)
    if stripped in DOMAIN_ALIAS_MAP:
        return DOMAIN_ALIAS_MAP[stripped]
    return domain_token.strip()


def source_column_to_dictionary_key(source_column: str, measure_type: str) -> str:
    upper = source_column.upper()
    summary_match = re.search(r"SUMMARY-RESULTS\[(?P<token>[A-Z0-9_]+)::SICC_(SYSTEM|TESTER)_AT_REF\]_CLASSHOT$", upper)
    if summary_match:
        token = summary_match.group("token")
        suffix = "SICC_SYSTEM_AT_REF" if measure_type == "system_at_ref" else "SICC_TESTER_AT_REF"
        return f"VA-NA-{token}-{suffix}_CLASSHOT"

    direct_match = re.search(r"VA-NA-(?P<token>[A-Z0-9_]+)-SICC_(SYSTEM|TESTER)_AT_REF_CLASSHOT$", upper)
    if direct_match:
        token = direct_match.group("token")
        suffix = "SICC_SYSTEM_AT_REF" if measure_type == "system_at_ref" else "SICC_TESTER_AT_REF"
        return f"VA-NA-{token}-{suffix}_CLASSHOT"

    return upper


def process_target_for_domain(canonical_domain: str) -> float:
    if canonical_domain in {"Atom", "Core", "Ring"}:
        return 0.95
    if canonical_domain == "GT":
        return 0.98
    return 0.9


def discover_measure_columns(columns: list[str]) -> dict[str, MeasureColumns]:
    discovered: dict[str, dict[str, str | None]] = {}
    for column in columns:
        upper = column.upper()
        token = extract_domain_token(upper)
        if not token:
            continue
        bucket = discovered.setdefault(token, {"raw": None, "system_at_ref": None, "tester_at_ref": None})
        if "SICC_SYSTEM_AT_REF" in upper:
            bucket["system_at_ref"] = column
        elif "SICC_TESTER_AT_REF" in upper:
            bucket["tester_at_ref"] = column
        elif "SUMMARY-RESULTS[" in upper and "::" not in upper:
            bucket["raw"] = column
        elif "RECLOGIN-" in upper:
            bucket["raw"] = column

    return {
        token: MeasureColumns(
            raw=values["raw"],
            system_at_ref=values["system_at_ref"],
            tester_at_ref=values["tester_at_ref"],
        )
        for token, values in discovered.items()
    }


def discover_true_raw_sicc_columns(columns: list[str]) -> list[tuple[str, str, int | None]]:
    raw_entries: list[tuple[str, str, int | None]] = []
    pattern = re.compile(r"^SICC-U1PU\d+_(?P<token>[A-Z0-9]+)-V1_value$", re.IGNORECASE)
    for column in columns:
        match = pattern.match(column)
        if not match:
            continue
        token = match.group("token").upper()
        core_number = None
        canonical_domain = token
        if token.startswith("IA") and len(token) >= 4 and token[2:].isdigit():
            canonical_domain = "Core"
            core_number = int(token[2:]) + 1
        elif token.startswith("AT") and len(token) >= 4 and token[2:].isdigit():
            canonical_domain = "Atom"
            core_number = int(token[2:]) + 1
        else:
            canonical_domain = RAW_DOMAIN_LABELS.get(token, token)
        raw_entries.append((column, canonical_domain, core_number))
    return raw_entries


def discover_upm_columns(columns: list[str]) -> dict[str, str]:
    upm_columns: dict[str, str] = {}
    for column in columns:
        if column == "UPM FEM CDIE":
            upm_columns["cpu"] = column
        elif column == "UPM FEM GT":
            upm_columns["gt"] = column
        elif column == "TPI_UPM::UPM_X_SCREEN_K_START_X_X_X_X_FEM_CALC_UPM_FEM_LVT_ATOMALL_PDK0P9_0950MV_SORT_U1.U2":
            upm_columns["other"] = column
    return upm_columns


def load_presi_targets(target_xlsx: Path) -> dict[str, float]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:
        return load_presi_targets_fallback(target_xlsx)

    workbook = load_workbook(target_xlsx, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header = [cell for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        domain_index = None
        target_index = None
        for index, value in enumerate(header):
            if isinstance(value, str) and value.strip().lower() == "domain":
                domain_index = index
            if isinstance(value, str) and value.strip().lower() == "sicc v1 presi target":
                target_index = index
        if domain_index is None or target_index is None:
            raise ValueError("Could not find 'Domain' and 'SICC V1 PreSi target' columns in the target workbook.")

        targets: dict[str, float] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            raw_domain = row[domain_index]
            raw_target = row[target_index]
            if raw_domain is None or raw_target is None:
                continue
            domain = canonical_domain_name(str(raw_domain))
            target = float(raw_target)
            if domain in targets and abs(targets[domain] - target) > 1e-9:
                raise ValueError(f"Multiple PreSi targets were found for domain '{domain}'.")
            targets[domain] = target
        return targets
    finally:
        workbook.close()


def load_presi_targets_fallback(target_xlsx: Path) -> dict[str, float]:
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

    def column_index_from_ref(cell_ref: str) -> int:
        letters = re.match(r"([A-Z]+)", cell_ref.upper())
        if not letters:
            raise ValueError(f"Invalid cell reference: {cell_ref}")
        value = 0
        for char in letters.group(1):
            value = (value * 26) + (ord(char) - ord("A") + 1)
        return value - 1

    def read_cell(cell: ET.Element, shared_strings: list[str]) -> str | None:
        cell_type = cell.attrib.get("t")
        if cell_type == "inlineStr":
            node = cell.find("x:is/x:t", namespace)
            return node.text if node is not None else None
        value_node = cell.find("x:v", namespace)
        if value_node is None or value_node.text is None:
            return None
        if cell_type == "s":
            return shared_strings[int(value_node.text)]
        return value_node.text

    with zipfile.ZipFile(target_xlsx) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("x:si", namespace):
                text_parts = [node.text or "" for node in item.findall(".//x:t", namespace)]
                shared_strings.append("".join(text_parts))

        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        first_sheet = workbook_root.find("x:sheets/x:sheet", namespace)
        if first_sheet is None:
            raise ValueError("The target workbook does not contain any sheets.")
        rel_id = first_sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        rel_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target_sheet = None
        for relation in rel_root.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
            if relation.attrib.get("Id") == rel_id:
                target_sheet = "xl/" + relation.attrib["Target"].lstrip("/")
                break
        if target_sheet is None:
            raise ValueError("Could not resolve the first worksheet path from the target workbook.")

        sheet_root = ET.fromstring(archive.read(target_sheet))
        rows: list[list[str | None]] = []
        for row in sheet_root.findall("x:sheetData/x:row", namespace):
            values: list[str | None] = []
            for cell in row.findall("x:c", namespace):
                index = column_index_from_ref(cell.attrib["r"])
                while len(values) <= index:
                    values.append(None)
                values[index] = read_cell(cell, shared_strings)
            rows.append(values)

    if not rows:
        raise ValueError("The target workbook is empty.")
    header = rows[0]
    domain_index = None
    target_index = None
    for index, value in enumerate(header):
        if isinstance(value, str) and value.strip().lower() == "domain":
            domain_index = index
        if isinstance(value, str) and value.strip().lower() == "sicc v1 presi target":
            target_index = index
    if domain_index is None or target_index is None:
        raise ValueError("Could not find 'Domain' and 'SICC V1 PreSi target' columns in the target workbook.")

    targets: dict[str, float] = {}
    for row in rows[1:]:
        raw_domain = row[domain_index] if domain_index < len(row) else None
        raw_target = row[target_index] if target_index < len(row) else None
        if raw_domain is None or raw_target is None:
            continue
        domain = canonical_domain_name(str(raw_domain))
        target = float(raw_target)
        if domain in targets and abs(targets[domain] - target) > 1e-9:
            raise ValueError(f"Multiple PreSi targets were found for domain '{domain}'.")
        targets[domain] = target
    return targets


def load_token_dictionary(token_xlsx: Path) -> dict[str, TokenRule]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:
        return load_token_dictionary_fallback(token_xlsx)

    workbook = load_workbook(token_xlsx, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    return parse_token_dictionary_rows(rows)


def load_token_dictionary_fallback(token_xlsx: Path) -> dict[str, TokenRule]:
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

    def column_index_from_ref(cell_ref: str) -> int:
        letters = re.match(r"([A-Z]+)", cell_ref.upper())
        if not letters:
            raise ValueError(f"Invalid cell reference: {cell_ref}")
        value = 0
        for char in letters.group(1):
            value = (value * 26) + (ord(char) - ord("A") + 1)
        return value - 1

    def read_cell(cell: ET.Element, shared_strings: list[str]) -> str | None:
        cell_type = cell.attrib.get("t")
        if cell_type == "inlineStr":
            node = cell.find("x:is/x:t", namespace)
            return node.text if node is not None else None
        value_node = cell.find("x:v", namespace)
        if value_node is None or value_node.text is None:
            return None
        if cell_type == "s":
            return shared_strings[int(value_node.text)]
        return value_node.text

    with zipfile.ZipFile(token_xlsx) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("x:si", namespace):
                text_parts = [node.text or "" for node in item.findall(".//x:t", namespace)]
                shared_strings.append("".join(text_parts))

        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        first_sheet = workbook_root.find("x:sheets/x:sheet", namespace)
        if first_sheet is None:
            raise ValueError("The token dictionary workbook does not contain any sheets.")
        rel_id = first_sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        rel_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target_sheet = None
        for relation in rel_root.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
            if relation.attrib.get("Id") == rel_id:
                target_sheet = "xl/" + relation.attrib["Target"].lstrip("/")
                break
        if target_sheet is None:
            raise ValueError("Could not resolve the first worksheet path from the token dictionary workbook.")

        sheet_root = ET.fromstring(archive.read(target_sheet))
        rows: list[list[str | None]] = []
        for row in sheet_root.findall("x:sheetData/x:row", namespace):
            values: list[str | None] = []
            for cell in row.findall("x:c", namespace):
                index = column_index_from_ref(cell.attrib["r"])
                while len(values) <= index:
                    values.append(None)
                values[index] = read_cell(cell, shared_strings)
            rows.append(values)

    return parse_token_dictionary_rows(rows)


def parse_token_dictionary_rows(rows: list[tuple[object, ...]] | list[list[object | None]]) -> dict[str, TokenRule]:
    token_rules: dict[str, TokenRule] = {}
    if not rows:
        return token_rules

    for row in rows[1:]:
        token = str(row[0]).strip() if len(row) >= 1 and row[0] is not None else ""
        canonical_domain = str(row[1]).strip() if len(row) >= 2 and row[1] is not None else ""
        core_text = str(row[2]).strip() if len(row) >= 3 and row[2] is not None else ""
        if not token or not canonical_domain:
            continue
        core_number = None
        if core_text:
            try:
                core_number = int(float(core_text))
            except ValueError:
                core_number = None
        token_rules[token.upper()] = TokenRule(
            source_token=token,
            canonical_domain=canonical_domain,
            core_number=core_number,
        )
    return token_rules


def parse_end_time(series: pd.Series) -> pd.Series:
    parsed = pd.to_datetime(series, errors="coerce", utc=True)
    return parsed.dt.tz_convert("UTC").dt.tz_localize(None)


def parse_collapsed_numeric_value(raw_value: object) -> tuple[float | None, str, int]:
    if raw_value is None:
        return None, "empty", 0
    if isinstance(raw_value, float) and math.isnan(raw_value):
        return None, "empty", 0
    text = str(raw_value).strip()
    if not text or text.lower() == "nan":
        return None, "empty", 0
    segments = [segment.strip() for segment in text.split("~") if segment.strip()]
    if not segments:
        return None, "empty", 0

    numeric_values: list[float] = []
    for segment in segments:
        try:
            numeric_values.append(float(segment))
        except ValueError:
            return None, "non_numeric", len(segments)

    if len(numeric_values) == 1:
        return numeric_values[0], "scalar", 1

    reference = numeric_values[0]
    tolerance = max(1e-9, abs(reference) * 1e-6)
    if max(abs(value - reference) for value in numeric_values) <= tolerance:
        return reference, "collapsed", len(segments)
    return None, "mixed_segments", len(segments)


def load_latest_classhot_rows(
    input_csv: Path,
    use_columns: list[str],
    visual_id_column: str,
    end_time_column: str,
    chunk_size: int,
) -> pd.DataFrame:
    best_rows: dict[str, dict[str, object]] = {}
    best_times: dict[str, pd.Timestamp] = {}

    for chunk in pd.read_csv(input_csv, usecols=use_columns, chunksize=chunk_size, low_memory=False):
        chunk[end_time_column] = parse_end_time(chunk[end_time_column])
        chunk = chunk.dropna(subset=[visual_id_column])
        chunk[visual_id_column] = chunk[visual_id_column].astype(str).str.strip()
        chunk = chunk[chunk[visual_id_column] != ""]
        for row in chunk.to_dict(orient="records"):
            visual_id = str(row[visual_id_column]).strip()
            end_time = row[end_time_column]
            previous_time = best_times.get(visual_id)
            if previous_time is None or (pd.notna(end_time) and (pd.isna(previous_time) or end_time >= previous_time)):
                best_times[visual_id] = end_time
                best_rows[visual_id] = row

    latest_rows = pd.DataFrame(best_rows.values())
    if not latest_rows.empty:
        latest_rows.sort_values([end_time_column, visual_id_column], ascending=[False, True], inplace=True, na_position="last")
    return latest_rows


def upm_bucket_for_domain(canonical_domain: str) -> str:
    if canonical_domain in {"Atom", "Core", "Ring"}:
        return "cpu"
    if canonical_domain == "GT":
        return "gt"
    return "other"


def resolve_raw_column_for_domain(canonical_domain: str, measure_columns: dict[str, MeasureColumns]) -> tuple[str | None, str | None]:
    raw_columns = {token: columns.raw for token, columns in measure_columns.items() if columns.raw}
    for raw_token in RAW_ALIAS_PRIORITY.get(canonical_domain, []):
        if raw_token in raw_columns:
            note = None
            if canonical_domain == "Ring":
                note = "Ring raw uses CORE_OUT because the input schema has no dedicated Ring raw column."
            return raw_columns[raw_token], note
    return None, None


def build_metric_configs(measure_columns: dict[str, MeasureColumns], upm_columns: dict[str, str], token_rules: dict[str, TokenRule]) -> list[MetricConfig]:
    configs: list[MetricConfig] = []

    for token, columns in sorted(measure_columns.items()):
        rule = None
        if columns.system_at_ref:
            rule = token_rules.get(source_column_to_dictionary_key(columns.system_at_ref, "system_at_ref"))
        if rule is None and columns.tester_at_ref:
            rule = token_rules.get(source_column_to_dictionary_key(columns.tester_at_ref, "tester_at_ref").replace("SICC_TESTER_AT_REF", "SICC_SYSTEM_AT_REF"))
        if rule is None:
            continue

        canonical_domain = canonical_domain_name(rule.canonical_domain)
        upm_column = upm_columns[upm_bucket_for_domain(canonical_domain)]

        if columns.system_at_ref:
            configs.append(
                MetricConfig(
                    canonical_domain=canonical_domain,
                    domain_token=token,
                    measure_type="system_at_ref",
                    source_column=columns.system_at_ref,
                    upm_column=upm_column,
                    raw_mapping_note=None,
                    core_number=rule.core_number,
                )
            )
        if columns.tester_at_ref:
            configs.append(
                MetricConfig(
                    canonical_domain=canonical_domain,
                    domain_token=token,
                    measure_type="tester_at_ref",
                    source_column=columns.tester_at_ref,
                    upm_column=upm_column,
                    raw_mapping_note=None,
                    core_number=rule.core_number,
                )
            )
    return configs


def build_raw_metric_configs(columns: list[str], upm_columns: dict[str, str]) -> list[MetricConfig]:
    configs: list[MetricConfig] = []
    for source_column, canonical_domain, core_number in discover_true_raw_sicc_columns(columns):
        upm_column = upm_columns[upm_bucket_for_domain(canonical_domain)]
        configs.append(
            MetricConfig(
                canonical_domain=canonical_domain,
                domain_token=source_column,
                measure_type="raw",
                source_column=source_column,
                upm_column=upm_column,
                raw_mapping_note=None,
                core_number=core_number,
            )
        )
    return configs


def build_normalized_points(
    latest_rows: pd.DataFrame,
    required_columns: dict[str, str],
    all_columns: list[str],
    measure_columns: dict[str, MeasureColumns],
    upm_columns: dict[str, str],
    targets: dict[str, float],
    token_rules: dict[str, TokenRule],
) -> tuple[pd.DataFrame, dict[str, object]]:
    configs = build_metric_configs(measure_columns, upm_columns, token_rules) + build_raw_metric_configs(all_columns, upm_columns)
    records: list[dict[str, object]] = []
    parse_audit: dict[str, dict[str, int]] = {}

    for config in configs:
        parse_audit.setdefault(config.source_column, {"scalar": 0, "collapsed": 0, "mixed_segments": 0, "non_numeric": 0, "empty": 0})
        upm_series = pd.to_numeric(latest_rows[config.upm_column], errors="coerce")
        source_series = latest_rows[config.source_column]
        for row_index, raw_value in source_series.items():
            parsed_value, parse_state, segment_count = parse_collapsed_numeric_value(raw_value)
            parse_audit[config.source_column][parse_state] += 1
            if parsed_value is None:
                continue

            upm_value = upm_series.loc[row_index]
            if pd.isna(upm_value):
                continue

            records.append(
                {
                    "product": str(latest_rows.at[row_index, required_columns["prod"]]).strip().upper(),
                    "visual_id": str(latest_rows.at[row_index, required_columns["visual_id"]]).strip(),
                    "program_name": str(latest_rows.at[row_index, required_columns["program_name"]]).strip(),
                    "end_time": latest_rows.at[row_index, required_columns["end_time"]],
                    "canonical_domain": config.canonical_domain,
                    "domain_token": config.domain_token,
                    "measure_type": config.measure_type,
                    "sicc_value": float(parsed_value),
                    "upm_value": float(upm_value),
                    "process_target": process_target_for_domain(config.canonical_domain),
                    "presi_target": targets.get(config.canonical_domain),
                    "core_number": config.core_number,
                    "source_column": config.source_column,
                    "upm_column": config.upm_column,
                    "parse_state": parse_state,
                    "segment_count": segment_count,
                    "raw_mapping_note": config.raw_mapping_note,
                }
            )

    points_df = pd.DataFrame.from_records(records)
    if not points_df.empty:
        points_df.sort_values(["canonical_domain", "measure_type", "product", "domain_token", "visual_id"], inplace=True)
    return points_df, parse_audit


def fit_log_curve(group_df: pd.DataFrame, process_target: float) -> dict[str, float] | None:
    x = group_df["upm_value"].to_numpy(dtype=float)
    y = group_df["sicc_value"].to_numpy(dtype=float)
    mask = np.isfinite(x) & np.isfinite(y) & (x > 0.0) & (y > 0.0)
    if int(mask.sum()) < 3:
        return None

    x = x[mask]
    y = y[mask]
    log_y = np.log(y)
    design = np.column_stack([np.ones(x.shape[0]), x])
    coefficients, _, _, _ = np.linalg.lstsq(design, log_y, rcond=None)
    predictions = design @ coefficients
    ss_res = float(np.sum((log_y - predictions) ** 2))
    ss_tot = float(np.sum((log_y - float(np.mean(log_y))) ** 2))
    r2 = 1.0 if ss_tot <= 1e-12 else 1.0 - (ss_res / ss_tot)
    intercept = float(coefficients[0])
    slope = float(coefficients[1])
    normalized_at_target = float(math.exp(intercept + (slope * process_target)))
    return {
        "n": int(x.shape[0]),
        "slope_linear": slope,
        "intercept_log": intercept,
        "r2": r2,
        "normalized_at_target": normalized_at_target,
        "x_min": float(np.min(x)),
        "x_max": float(np.max(x)),
    }


def relative_spread(values: list[float]) -> float:
    if not values:
        return 0.0
    mean_value = float(np.mean(values))
    if abs(mean_value) <= 1e-12:
        return 0.0
    return (max(values) - min(values)) / abs(mean_value)


def select_points_and_build_comments(points_df: pd.DataFrame, core_shift_threshold_pct: float) -> tuple[pd.DataFrame, pd.DataFrame]:
    if points_df.empty:
        return points_df.copy(), pd.DataFrame()

    threshold = core_shift_threshold_pct / 100.0
    selected_frames: list[pd.DataFrame] = []
    comments: list[dict[str, object]] = []

    for (canonical_domain, measure_type), group_df in points_df.groupby(["canonical_domain", "measure_type"], sort=True):
        token_order_rows = (
            group_df[["domain_token", "core_number"]]
            .drop_duplicates()
            .sort_values(["core_number", "domain_token"], na_position="last")
        )
        tokens = token_order_rows["domain_token"].astype(str).tolist()
        if len(tokens) <= 1:
            retained = group_df.copy()
            retained["selection_mode"] = "single_token"
            selected_frames.append(retained)
            continue

        process_target = float(group_df["process_target"].iloc[0])
        significant_shift = False
        for product, product_df in group_df.groupby("product", sort=True):
            token_values: list[float] = []
            for token in tokens:
                fit = fit_log_curve(product_df[product_df["domain_token"] == token], process_target)
                if fit is not None:
                    token_values.append(fit["normalized_at_target"])
            if len(token_values) < 2:
                continue
            spread = relative_spread(token_values)
            if spread > threshold:
                significant_shift = True
                comments.append(
                    {
                        "comment_type": "core_shift",
                        "canonical_domain": canonical_domain,
                        "measure_type": measure_type,
                        "product": product,
                        "message": f"{canonical_domain} {measure_type}: token spread is {spread * 100.0:.2f}% across {', '.join(tokens)}.",
                    }
                )

        if significant_shift:
            retained = group_df.copy()
            retained["selection_mode"] = "retain_all_tokens"
            selected_frames.append(retained)
        else:
            retained = group_df[group_df["domain_token"] == tokens[0]].copy()
            retained["selection_mode"] = "first_token_only"
            selected_frames.append(retained)

    selected_points = pd.concat(selected_frames, ignore_index=True) if selected_frames else pd.DataFrame()
    comments_df = pd.DataFrame.from_records(comments)
    return selected_points, comments_df


def add_program_shift_comments(selected_points: pd.DataFrame, comments_df: pd.DataFrame, threshold_pct: float) -> pd.DataFrame:
    comments = comments_df.to_dict(orient="records") if not comments_df.empty else []
    threshold = threshold_pct / 100.0
    for (canonical_domain, measure_type, product), group_df in selected_points.groupby(["canonical_domain", "measure_type", "product"], sort=True):
        process_target = float(group_df["process_target"].iloc[0])
        per_program_values: list[float] = []
        program_names: list[str] = []
        for program_name, program_df in group_df.groupby("program_name", sort=True):
            fit = fit_log_curve(program_df, process_target)
            if fit is None or fit["n"] < 8:
                continue
            per_program_values.append(fit["normalized_at_target"])
            program_names.append(str(program_name))
        if len(per_program_values) < 2:
            continue
        spread = relative_spread(per_program_values)
        if spread > threshold:
            comments.append(
                {
                    "comment_type": "program_shift",
                    "canonical_domain": canonical_domain,
                    "measure_type": measure_type,
                    "product": product,
                    "message": f"{canonical_domain} {measure_type}: program spread is {spread * 100.0:.2f}% across {', '.join(program_names)}.",
                }
            )
    return pd.DataFrame.from_records(comments)


def apply_program_scope(selected_points: pd.DataFrame, comments_df: pd.DataFrame, threshold_pct: float) -> tuple[pd.DataFrame, pd.DataFrame]:
    comments = comments_df.to_dict(orient="records") if not comments_df.empty else []
    threshold = threshold_pct / 100.0
    filtered_frames: list[pd.DataFrame] = []

    for (canonical_domain, measure_type, product), group_df in selected_points.groupby(["canonical_domain", "measure_type", "product"], sort=True):
        process_target = float(group_df["process_target"].iloc[0])
        program_rows: list[dict[str, object]] = []
        for program_name, program_df in group_df.groupby("program_name", sort=True):
            fit = fit_log_curve(program_df, process_target)
            if fit is None or fit["n"] < 20:
                continue
            latest_program_time = pd.to_datetime(program_df["end_time"], errors="coerce").max()
            program_rows.append(
                {
                    "program_name": str(program_name),
                    "normalized_at_target": fit["normalized_at_target"],
                    "count": fit["n"],
                    "latest_time": latest_program_time,
                }
            )

        if len(program_rows) < 2:
            filtered_frames.append(group_df.copy())
            continue

        spread = relative_spread([float(row["normalized_at_target"]) for row in program_rows])
        if spread <= threshold:
            filtered_frames.append(group_df.copy())
            continue

        sorted_programs = sorted(
            program_rows,
            key=lambda row: (
                pd.Timestamp.min if pd.isna(row["latest_time"]) else row["latest_time"],
                row["count"],
                row["program_name"],
            ),
            reverse=True,
        )
        chosen_program = str(sorted_programs[0]["program_name"])
        filtered_group = group_df[group_df["program_name"] == chosen_program].copy()
        filtered_group["selection_mode"] = filtered_group["selection_mode"].astype(str) + "|latest_program_only"
        filtered_frames.append(filtered_group)
        comments.append(
            {
                "comment_type": "program_scope",
                "canonical_domain": canonical_domain,
                "measure_type": measure_type,
                "product": product,
                "message": f"{canonical_domain} {measure_type}: using latest significant-shift program {chosen_program} for normalized summary and plot because program spread is {spread * 100.0:.2f}%.",
            }
        )

    filtered_points = pd.concat(filtered_frames, ignore_index=True) if filtered_frames else pd.DataFrame()
    return filtered_points, pd.DataFrame.from_records(comments)


def build_summary_table(selected_points: pd.DataFrame, program_shift_threshold_pct: float) -> pd.DataFrame:
    summary_rows: list[dict[str, object]] = []
    for (canonical_domain, measure_type, product), group_df in selected_points.groupby(["canonical_domain", "measure_type", "product"], sort=True):
        process_target = float(group_df["process_target"].iloc[0])
        fit_group = resolve_fit_points(group_df, program_shift_threshold_pct)
        fit = fit_log_curve(fit_group, process_target)
        if fit is None:
            continue
        summary_rows.append(
            {
                "canonical_domain": canonical_domain,
                "measure_type": measure_type,
                "product": product,
                "selected_tokens": ", ".join(sorted(group_df["domain_token"].dropna().astype(str).unique().tolist())),
                "selection_mode": ", ".join(sorted(group_df["selection_mode"].dropna().astype(str).unique().tolist())),
                "upm_column": str(group_df["upm_column"].iloc[0]),
                "process_target": process_target,
                "presi_target": group_df["presi_target"].iloc[0],
                "point_count": fit["n"],
                "visual_count": int(fit_group["visual_id"].nunique()),
                "program_count": int(fit_group["program_name"].astype(str).str.strip().replace("", np.nan).dropna().nunique()),
                "normalized_sicc_at_target": fit["normalized_at_target"],
                "slope_linear": fit["slope_linear"],
                "intercept_log": fit["intercept_log"],
                "r2": fit["r2"],
            }
        )
    summary_df = pd.DataFrame.from_records(summary_rows)
    if not summary_df.empty:
        summary_df.sort_values(["canonical_domain", "measure_type", "product"], inplace=True)
    return summary_df


def domain_sort_key(domain_name: str) -> tuple[int, str]:
    preferred = ["Core", "Ring", "Atom", "SA-Atom", "GT", "SAHUB", "NPU", "Hub Atom", "SACD", "SAME", "SA"]
    if domain_name in preferred:
        return (preferred.index(domain_name), domain_name)
    return (len(preferred), domain_name)


def build_wide_summary_table(summary_df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    products = ["NVLAX", "NVLS7", "NVLHX"]
    measures = ["raw", "tester_at_ref", "system_at_ref"]

    if summary_df.empty:
        return pd.DataFrame(columns=["domain"] + [f"{measure}__{product}" for measure in measures for product in products])

    for domain_name in sorted(summary_df["canonical_domain"].dropna().astype(str).unique().tolist(), key=domain_sort_key):
        row: dict[str, object] = {"domain": domain_name}
        for measure_type in measures:
            for product in products:
                match = summary_df[
                    (summary_df["canonical_domain"] == domain_name)
                    & (summary_df["measure_type"] == measure_type)
                    & (summary_df["product"] == product)
                ]
                value = None
                if not match.empty:
                    value = round(float(match.iloc[0]["normalized_sicc_at_target"]), 4)
                row[f"{measure_type}__{product}"] = value
        rows.append(row)

    return pd.DataFrame.from_records(rows)


def build_overview_table(selected_points: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    if selected_points.empty:
        return pd.DataFrame(columns=["product", "sample_size", "median_fem", "relevant_test_programs"])

    for product, group_df in selected_points.groupby("product", sort=True):
        program_counts = group_df["program_name"].astype(str).value_counts()
        rows.append(
            {
                "product": product,
                "sample_size": int(group_df.shape[0]),
                "median_fem": float(pd.to_numeric(group_df["upm_value"], errors="coerce").median()),
                "relevant_test_programs": "; ".join(f"{name} ({count})" for name, count in program_counts.items()),
            }
        )
    return pd.DataFrame.from_records(rows)


def trim_sigma_outliers(points_df: pd.DataFrame, sigma_multiplier: float = 7.0) -> pd.DataFrame:
    if points_df.empty:
        return points_df.copy()

    filtered_frames: list[pd.DataFrame] = []
    for (canonical_domain, measure_type, product), group_df in points_df.groupby(["canonical_domain", "measure_type", "product"], sort=True):
        sicc_values = pd.to_numeric(group_df["sicc_value"], errors="coerce")
        valid = np.isfinite(sicc_values)
        if valid.sum() < 25:
            filtered_frames.append(group_df.copy())
            continue

        median_value = float(np.nanmedian(sicc_values[valid].to_numpy()))
        sigma_value = float(np.nanstd(sicc_values[valid].to_numpy()))
        if not np.isfinite(median_value) or not np.isfinite(sigma_value) or sigma_value <= 0.0:
            filtered_frames.append(group_df.copy())
            continue

        lower_bound = median_value - (sigma_multiplier * sigma_value)
        upper_bound = median_value + (sigma_multiplier * sigma_value)
        retained = group_df[valid & (sicc_values >= lower_bound) & (sicc_values <= upper_bound)].copy()
        if retained.empty:
            retained = group_df.copy()
        filtered_frames.append(retained)

    return pd.concat(filtered_frames, ignore_index=True) if filtered_frames else points_df.copy()


def resolve_fit_points(group_df: pd.DataFrame, threshold_pct: float) -> pd.DataFrame:
    token_order_rows = (
        group_df[["domain_token", "core_number"]]
        .drop_duplicates()
        .sort_values(["core_number", "domain_token"], na_position="last")
    )
    tokens = token_order_rows["domain_token"].astype(str).tolist()
    if tokens:
        group_df = group_df[group_df["domain_token"] == tokens[0]].copy()

    fit_points, _ = apply_program_scope(group_df, pd.DataFrame(), threshold_pct)
    return fit_points if not fit_points.empty else group_df.copy()


def add_summary_overview_slide(presentation: Presentation, overview_df: pd.DataFrame) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.25), Inches(0.12), Inches(12.8), Inches(0.45))
    title_frame = title_box.text_frame
    title_frame.text = "Analysis Overview"
    title_frame.paragraphs[0].font.size = Inches(0.26)
    title_frame.paragraphs[0].font.bold = True

    subtitle_box = slide.shapes.add_textbox(Inches(0.25), Inches(0.52), Inches(12.8), Inches(0.3))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "All products and test programs retained; only units beyond +/- 7 sigma from the median are removed."
    subtitle_frame.paragraphs[0].font.size = Inches(0.14)

    rows = len(overview_df.index) + 1
    cols = 4
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.25), Inches(0.9), Inches(12.8), Inches(6.0))
    table = table_shape.table
    table.columns[0].width = Inches(1.1)
    table.columns[1].width = Inches(1.3)
    table.columns[2].width = Inches(1.2)
    table.columns[3].width = Inches(9.2)

    headers = ["Product", "Sample size", "Median FEM", "Relevant test programs"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header

    for row_index, (_, row) in enumerate(overview_df.iterrows(), start=1):
        table.cell(row_index, 0).text = str(row["product"])
        table.cell(row_index, 1).text = f"{int(row['sample_size'])}"
        table.cell(row_index, 2).text = f"{float(row['median_fem']):.4f}"
        table.cell(row_index, 3).text = str(row["relevant_test_programs"])

    for row_index in range(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = RGBColor(219, 234, 254)
            elif row_index % 2 == 1:
                cell.fill.fore_color.rgb = RGBColor(248, 250, 252)
            else:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Inches(0.14)
                if row_index == 0:
                    paragraph.font.bold = True



def add_plot_slide(presentation: Presentation, image_path: Path) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_width = 13.333
    slide_height = 7.5
    left_margin = 0.08
    top_margin = 0.06
    right_margin = 0.08
    bottom_margin = 0.06
    picture_width = slide_width - left_margin - right_margin
    picture_height = slide_height - top_margin - bottom_margin
    slide.shapes.add_picture(str(image_path), Inches(left_margin), Inches(top_margin), width=Inches(picture_width), height=Inches(picture_height))


def trim_plot_focus_window(group_df: pd.DataFrame) -> pd.DataFrame:
    if group_df.empty:
        return group_df.copy()

    sicc_values = pd.to_numeric(group_df["sicc_value"], errors="coerce")
    valid = np.isfinite(sicc_values)
    positive_values = sicc_values[valid]
    if positive_values.shape[0] < 25:
        return group_df.copy()

    median_value = float(np.nanmedian(positive_values.to_numpy()))
    sigma_value = float(np.nanstd(positive_values.to_numpy()))
    if not np.isfinite(median_value) or not np.isfinite(sigma_value) or sigma_value <= 0.0:
        return group_df.copy()

    lower_bound = median_value - (7.0 * sigma_value)
    upper_bound = median_value + (7.0 * sigma_value)
    focus_df = group_df[valid & (sicc_values >= lower_bound) & (sicc_values <= upper_bound)].copy()
    if focus_df.empty:
        return group_df.copy()
    return focus_df


def plot_domain_measure(group_df: pd.DataFrame, comments_df: pd.DataFrame, output_path: Path, program_shift_threshold_pct: float) -> dict[str, object] | None:
    if group_df.empty:
        return None

    canonical_domain = str(group_df["canonical_domain"].iloc[0])
    measure_type = str(group_df["measure_type"].iloc[0])
    process_target = float(group_df["process_target"].iloc[0])
    presi_target = group_df["presi_target"].iloc[0]
    fit_df = resolve_fit_points(group_df, program_shift_threshold_pct)
    plot_df = trim_plot_focus_window(group_df)
    x_min = float(plot_df["upm_value"].min())
    x_max = float(plot_df["upm_value"].max())
    x_pad = max(0.005, (x_max - x_min) * 0.05)
    plot_x_min = x_min - x_pad
    plot_x_max = x_max + x_pad
    y_min = float(plot_df["sicc_value"].min())
    y_max = float(plot_df["sicc_value"].max())
    y_pad_factor = 1.15
    plot_y_min = max(1e-6, y_min / y_pad_factor)
    plot_y_max = y_max * y_pad_factor

    fig, ax = plt.subplots(figsize=(13.6, 7.6), constrained_layout=True)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#fbfbfa")
    ax.grid(True, which="major", axis="both", color="#dde3ea", linewidth=0.8)
    ax.grid(True, which="minor", axis="y", color="#eef2f7", linewidth=0.5)
    ax.set_axisbelow(True)

    fit_notes: list[tuple[str, float, str]] = []
    line_x = np.linspace(max(1e-6, plot_x_min), plot_x_max, 400)
    for product in sorted(plot_df["product"].dropna().astype(str).unique().tolist()):
        product_df = plot_df[plot_df["product"] == product]
        color = PRODUCT_COLORS.get(product, "#444444")
        ax.scatter(product_df["upm_value"], product_df["sicc_value"], s=7, marker="s", alpha=0.22, color=color, edgecolors="none", label=PRODUCT_DISPLAY_NAMES.get(product, product))
        fit_source = fit_df[fit_df["product"] == product]
        fit = fit_log_curve(fit_source, process_target)
        if fit is None:
            continue
        fit_y = np.exp(fit["intercept_log"] + (fit["slope_linear"] * line_x))
        ax.plot(line_x, fit_y, color=color, linewidth=1.9)
        normalized_y = fit["normalized_at_target"]
        if normalized_y > 0:
            ax.axhline(normalized_y, color=color, linestyle="--", linewidth=0.9, alpha=0.9)
            fit_notes.append((product, normalized_y, color))

    ax.axvline(process_target, color="#111827", linestyle="--", linewidth=1.8)
    ax.text(process_target, 1.02, f"Process target={process_target:.2f}", transform=ax.get_xaxis_transform(), ha="center", va="bottom", fontsize=12.5, color="#111827")

    if pd.notna(presi_target) and float(presi_target) > 0.0:
        presi_value = float(presi_target)
        ax.axhline(presi_value, color="#6b7280", linestyle=":", linewidth=1.6)
        ax.text(plot_x_min + 0.01, presi_value * 1.01, f"PreSi target ={presi_value:.3f}A", fontsize=12, color="#4b5563", ha="left", va="bottom")

    measure_label = {
        "raw": "Raw SICC",
        "system_at_ref": "SICC System At Ref",
        "tester_at_ref": "SICC Tester At Ref",
    }.get(measure_type, measure_type)
    ax.set_title(f"{canonical_domain} {measure_label}", fontsize=18, fontweight="bold", pad=14)
    if fit_notes:
        label_x = plot_x_min + 0.01
        label_offsets = [0.03, 0.0, -0.03]
        for index, (product, normalized_y, color) in enumerate(sorted(fit_notes, key=lambda item: item[1], reverse=True)):
            y_label = normalized_y + (label_offsets[index] if index < len(label_offsets) else 0.0)
            label = f"{PRODUCT_DISPLAY_NAMES.get(product, product)} PostSi={normalized_y:.3f}" if product == "NVLAX" else f"{PRODUCT_DISPLAY_NAMES.get(product, product)} PostSi={normalized_y:.2f}"
            ax.text(label_x, y_label, label, color=color, fontsize=11, ha="left", va="center")

    ax.set_xlabel("UPM FEM", fontsize=12)
    ax.set_ylabel("SICC", fontsize=12)
    ax.set_yscale("log")
    ax.set_xlim(plot_x_min, plot_x_max)
    ax.set_ylim(plot_y_min, plot_y_max)
    if plot_y_min >= 1.0 and plot_y_max < 10.0:
        ax.yaxis.set_major_locator(LogLocator(base=10.0, subs=(1.0, 2.0, 5.0)))
        ax.yaxis.set_major_formatter(LogFormatter(base=10.0, labelOnlyBase=False))
        ax.yaxis.set_minor_formatter(NullFormatter())
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)

    fig.savefig(output_path, dpi=220, bbox_inches="tight")
    plt.close(fig)
    return {
        "canonical_domain": canonical_domain,
        "measure_type": measure_type,
        "image_path": str(output_path),
        "point_count": int(group_df.shape[0]),
        "visual_count": int(group_df["visual_id"].nunique()),
    }


def add_summary_table_slide(presentation: Presentation, wide_summary_df: pd.DataFrame) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.2), Inches(0.1), Inches(12.8), Inches(0.4))
    title_frame = title_box.text_frame
    title_frame.text = "Normalized SICC Summary"
    title_frame.paragraphs[0].font.size = Inches(0.24)
    title_frame.paragraphs[0].font.bold = True

    rows = len(wide_summary_df.index) + 2
    cols = 10
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.2), Inches(0.65), Inches(12.8), Inches(6.4))
    table = table_shape.table
    table.columns[0].width = Inches(1.55)
    for index in range(1, cols):
        table.columns[index].width = Inches(1.25)

    table.cell(0, 0).text = ""
    table.cell(0, 1).text = "raw"
    table.cell(0, 1).merge(table.cell(0, 3))
    table.cell(0, 4).text = "tester at ref"
    table.cell(0, 4).merge(table.cell(0, 6))
    table.cell(0, 7).text = "system at ref"
    table.cell(0, 7).merge(table.cell(0, 9))

    table.cell(1, 0).text = "domain"
    headers = ["NVLAX", "NVLS", "NVLHX", "NVLAX", "NVLS", "NVLHX", "NVLAX", "NVLS", "NVLHX"]
    for index, header in enumerate(headers, start=1):
        table.cell(1, index).text = header

    for row_index, (_, row) in enumerate(wide_summary_df.iterrows(), start=2):
        table.cell(row_index, 0).text = str(row["domain"])
        ordered_keys = [
            "raw__NVLAX", "raw__NVLS7", "raw__NVLHX",
            "tester_at_ref__NVLAX", "tester_at_ref__NVLS7", "tester_at_ref__NVLHX",
            "system_at_ref__NVLAX", "system_at_ref__NVLS7", "system_at_ref__NVLHX",
        ]
        for column_index, key in enumerate(ordered_keys, start=1):
            value = row.get(key)
            table.cell(row_index, column_index).text = "" if pd.isna(value) else f"{float(value):.4f}"

    for row_index in range(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = RGBColor(219, 234, 254)
            elif row_index == 1:
                cell.fill.fore_color.rgb = RGBColor(239, 246, 255)
            elif row_index % 2 == 0:
                cell.fill.fore_color.rgb = RGBColor(248, 250, 252)
            else:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if col_index else PP_ALIGN.LEFT
                if row_index <= 1:
                    paragraph.font.bold = True
                paragraph.font.size = Inches(0.14)


def build_powerpoint(deck_path: Path, plot_rows: list[dict[str, object]], overview_df: pd.DataFrame, wide_summary_df: pd.DataFrame) -> Path:
    deck_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    add_summary_overview_slide(presentation, overview_df)
    for row in sorted(plot_rows, key=lambda row: (str(row["canonical_domain"]), str(row["measure_type"]))):
        image_path = Path(str(row["image_path"]))
        add_plot_slide(presentation, image_path)
    add_summary_table_slide(presentation, wide_summary_df)
    try:
        presentation.save(str(deck_path))
        return deck_path
    except PermissionError:
        fallback = deck_path.with_name(f"{deck_path.stem}_{dt.datetime.now().strftime('%Y%m%d_%H%M%S')}{deck_path.suffix}")
        presentation.save(str(fallback))
        return fallback


def write_json(path: Path, payload: object) -> None:
    path.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")


def build_schema_manifest(required_columns: dict[str, str], measure_columns: dict[str, MeasureColumns], upm_columns: dict[str, str], targets: dict[str, float]) -> dict[str, object]:
    rows = []
    for token, columns in sorted(measure_columns.items()):
        canonical = canonical_domain_name(token)
        rows.append(
            {
                "domain_token": token,
                "canonical_domain": canonical,
                "raw_column": columns.raw,
                "system_at_ref_column": columns.system_at_ref,
                "tester_at_ref_column": columns.tester_at_ref,
                "presi_target": targets.get(canonical),
            }
        )
    return {"required_columns": required_columns, "upm_columns": upm_columns, "domain_mappings": rows}


def main() -> int:
    args = parse_args()
    input_csv = Path(args.input_csv)
    target_xlsx = Path(args.presi_target_xlsx)
    token_xlsx = Path(args.token_dictionary_xlsx)
    output_dir = Path(args.output_dir)
    deck_path = Path(args.deck_path)
    plots_dir = output_dir / "plots"
    output_dir.mkdir(parents=True, exist_ok=True)
    plots_dir.mkdir(parents=True, exist_ok=True)
    for child in plots_dir.iterdir():
        if child.is_file():
            child.unlink()

    header = read_csv_header(input_csv)
    required_columns = resolve_required_columns(header)
    measure_columns = discover_measure_columns(header)
    upm_columns = discover_upm_columns(header)
    targets = load_presi_targets(target_xlsx)
    token_rules = load_token_dictionary(token_xlsx)
    if not {"cpu", "gt", "other"}.issubset(upm_columns):
        missing = sorted({"cpu", "gt", "other"}.difference(upm_columns))
        raise ValueError(f"Missing required UPM axis columns: {', '.join(missing)}")

    selected_columns = {required_columns["prod"], required_columns["visual_id"], required_columns["program_name"], required_columns["end_time"], *upm_columns.values()}
    for raw_column, _, _ in discover_true_raw_sicc_columns(header):
        selected_columns.add(raw_column)
    for columns in measure_columns.values():
        for column in (columns.raw, columns.system_at_ref, columns.tester_at_ref):
            if column:
                selected_columns.add(column)

    latest_rows = load_latest_classhot_rows(input_csv, sorted(selected_columns), required_columns["visual_id"], required_columns["end_time"], args.chunk_size)
    write_json(output_dir / "schema_manifest.json", build_schema_manifest(required_columns, measure_columns, upm_columns, targets))
    write_json(output_dir / "presi_targets.json", targets)
    latest_rows.to_csv(output_dir / "latest_classhot_rows.csv", index=False)
    dedupe_summary = {
        "input_csv": str(input_csv),
        "dedupe_key": required_columns["visual_id"],
        "latest_timestamp_column": required_columns["end_time"],
        "unique_visual_ids": int(latest_rows[required_columns["visual_id"]].nunique()) if not latest_rows.empty else 0,
        "retained_rows": int(latest_rows.shape[0]),
        "products": sorted(latest_rows[required_columns["prod"]].dropna().astype(str).str.strip().str.upper().unique().tolist()) if not latest_rows.empty else [],
    }
    write_json(output_dir / "dedupe_summary.json", dedupe_summary)

    print(f"Input CSV: {input_csv}")
    print(f"Output directory: {output_dir}")
    print(f"Unique Visual IDs retained: {dedupe_summary['unique_visual_ids']}")
    print(f"Products: {', '.join(dedupe_summary['products'])}")
    if args.profile_only:
        print(f"Schema manifest: {output_dir / 'schema_manifest.json'}")
        print(f"Latest rows CSV: {output_dir / 'latest_classhot_rows.csv'}")
        return 0

    normalized_points, parse_audit = build_normalized_points(latest_rows, required_columns, header, measure_columns, upm_columns, targets, token_rules)
    normalized_points.to_csv(output_dir / "normalized_sicc_points.csv", index=False)
    write_json(output_dir / "sicc_parse_audit.json", parse_audit)

    selected_points = normalized_points.copy()
    selected_points["selection_mode"] = "all_points"
    comments_df = pd.DataFrame(columns=["comment_type", "canonical_domain", "measure_type", "product", "message"])
    selected_points = trim_sigma_outliers(selected_points)
    selected_points.to_csv(output_dir / "selected_sicc_points.csv", index=False)
    comments_df.to_csv(output_dir / "comments.csv", index=False)

    overview_df = build_overview_table(selected_points)
    summary_df = build_summary_table(selected_points, args.program_shift_threshold_pct)
    summary_df.to_csv(output_dir / "normalized_sicc_summary.csv", index=False)
    wide_summary_df = build_wide_summary_table(summary_df)
    wide_summary_df.to_csv(output_dir / "normalized_sicc_summary_wide.csv", index=False)

    plot_rows: list[dict[str, object]] = []
    plot_images: list[Path] = []
    for (canonical_domain, measure_type), group_df in selected_points.groupby(["canonical_domain", "measure_type"], sort=True):
        plot_path = plots_dir / f"{safe_name(canonical_domain)}_{safe_name(measure_type)}.png"
        plot_result = plot_domain_measure(group_df, comments_df, plot_path, args.program_shift_threshold_pct)
        if plot_result is not None:
            plot_rows.append(plot_result)
            plot_images.append(plot_path)
    pd.DataFrame.from_records(plot_rows).to_csv(output_dir / "plot_index.csv", index=False)

    build_powerpoint(deck_path, plot_rows, overview_df, wide_summary_df)
    run_manifest = {
        "input_csv": str(input_csv),
        "target_xlsx": str(target_xlsx),
        "output_dir": str(output_dir),
        "deck_path": str(deck_path),
        "retained_visual_ids": dedupe_summary["unique_visual_ids"],
        "normalized_point_count": int(normalized_points.shape[0]),
        "selected_point_count": int(selected_points.shape[0]),
        "plot_count": len(plot_images),
        "table_row_count": int(wide_summary_df.shape[0]),
    }
    write_json(output_dir / "run_manifest.json", run_manifest)

    print(f"Normalized points CSV: {output_dir / 'normalized_sicc_points.csv'}")
    print(f"Selected points CSV: {output_dir / 'selected_sicc_points.csv'}")
    print(f"Summary CSV: {output_dir / 'normalized_sicc_summary.csv'}")
    print(f"Wide summary CSV: {output_dir / 'normalized_sicc_summary_wide.csv'}")
    print(f"Comments CSV: {output_dir / 'comments.csv'}")
    print(f"PowerPoint deck: {deck_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
