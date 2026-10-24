from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

import colorsys

import matplotlib
from pptx import Presentation
from pptx.util import Inches

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.patches import Patch


DOMAIN_ALIASES = {
    "AT": "Atom",
    "ATOM": "Atom",
    "CR": "Core",
    "CORE": "Core",
    "CCF": "Ring",
    "CLR": "Ring",
    "RING": "Ring",
    "GT": "GT",
    "GTVG": "GTVG",
    "SAAT": "SA-Atom",
    "SAQ": "SAQ",
    "SAC": "SAC",
    "SACD": "SACD",
    "SADPU": "SADPU",
    "SAIOC": "SAIOC",
    "SAME": "SAME",
    "SAN": "SAN",
    "SAPS": "SAPS",
    "NPU": "NPU",
    "HUB": "HUB",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Render NVLAX DTS distribution charts with Matplotlib.")
    parser.add_argument(
        "--input-csv",
        default=r"R:\Products\NVL\NVL-AX\Weekly data pull\Vmin_20261010.csv",
        help="Merged weekly CSV containing ILAS_* DTS columns.",
    )
    parser.add_argument(
        "--lot",
        default="",
        help="Optional lot to filter. If omitted, the script uses the dominant lot when more than one lot is present.",
    )
    parser.add_argument(
        "--program",
        default="",
        help="Optional exact test program name to filter. If supplied, only rows matching that program are included.",
    )
    parser.add_argument(
        "--output-dir",
        default="",
        help="Directory for PNG and CSV outputs. Defaults next to the input CSV.",
    )
    return parser.parse_args()


def safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", value.strip())
    cleaned = re.sub(r"(^-+|-+$)", "", cleaned)
    return cleaned.lower() or "unknown"


def get_first_existing_column(candidate_names: list[str], available_names: list[str]) -> str | None:
    for candidate in candidate_names:
        if candidate in available_names:
            return candidate
    return None


def get_display_frequency_label(frequency_ghz: float) -> str:
    return f"{frequency_ghz:0.1f}"


def resolve_domain_name(domain_token: str) -> str:
    token = domain_token.strip().upper()
    base_token = re.sub(r"\d+$", "", token)
    return DOMAIN_ALIASES.get(base_token, base_token)


def module_token_from_text(setter: object, limiter_pattern: object) -> str:
    for source in (setter, limiter_pattern):
        if source is None:
            continue
        upper = str(source).strip().upper()
        if not upper:
            continue

        exact = re.search(r"(?P<module>(ARR|FUN|SCN)_[A-Z0-9]+(?:_[A-Z0-9]+)?)", upper)
        if exact:
            return exact.group("module")

        if re.search(r"(^|[^A-Z])(ARR|ARRAY)([^A-Z]|$)", upper):
            return "ARR"
        if re.search(r"(^|[^A-Z])(FUN|FUNCTIONAL|SBFT)([^A-Z]|$)", upper):
            return "FUN"
        if re.search(r"(^|[^A-Z])(SCN|SCAN|ATSPEED)([^A-Z]|$)", upper):
            return "SCN"

    return "UNKNOWN"


def module_sort_order(module_name: str) -> int:
    upper = module_name.upper()
    if upper.startswith("ARR"):
        return 1
    if upper.startswith("FUN"):
        return 2
    if upper.startswith("SCN"):
        return 3
    if upper == "UNKNOWN":
        return 9
    return 5


def corner_order_key(corner_name: str) -> tuple[int, int, str]:
    value = str(corner_name).upper()
    if value.startswith("FMIN"):
        return (0, 0, value)
    match = re.match(r"F(\d+)", value)
    if match:
        return (1, int(match.group(1)), value)
    if value.startswith("FMAX"):
        return (2, 999, value)
    return (3, 0, value)


def adjust_lightness(color: tuple[float, float, float] | tuple[float, float, float, float], delta: float) -> tuple[float, float, float] | tuple[float, float, float, float]:
    rgb = color[:3]
    r, g, b = rgb
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = min(max(l + delta, 0.0), 1.0)
    new_rgb = colorsys.hls_to_rgb(h, l, s)
    if len(color) == 4:
        return (*new_rgb, color[3])
    return new_rgb


def corner_color_map(corner_names: list[str]) -> dict[str, tuple[float, float, float] | tuple[float, float, float, float]]:
    palette = plt.get_cmap("tab10")
    ordered_corners = sorted(set(corner_names), key=corner_order_key)
    mapping: dict[str, tuple[float, float, float]] = {}
    for index, corner in enumerate(ordered_corners):
        mapping[str(corner)] = tuple(palette(index % 10))
    return mapping


def module_color_for_corner(corner_name: str, module_name: str, base_colors: dict[str, tuple[float, float, float]]) -> tuple[float, float, float]:
    base = base_colors.get(str(corner_name), tuple(plt.get_cmap("tab10")(0)))
    upper_module = str(module_name).upper()
    if upper_module.startswith("ARR"):
        return adjust_lightness(base, -0.22)
    if upper_module.startswith("FUN"):
        return adjust_lightness(base, 0.0)
    if upper_module.startswith("SCN"):
        return adjust_lightness(base, 0.18)
    return adjust_lightness(base, 0.12)


def quantile_linear(values: np.ndarray, quantile: float) -> float:
    if values.size == 1:
        return float(values[0])
    return float(np.quantile(values, quantile, method="linear"))


def build_observations(input_csv: Path, selected_lot: str, selected_program: str = "") -> tuple[pd.DataFrame, str, str | None, str, list[str]]:
    frame = pd.read_csv(input_csv, low_memory=False)
    columns = list(frame.columns)

    visual_id_column = get_first_existing_column(["VISUAL_ID", "Visual ID", "VisualId", "VISUALID", "VID"], columns)
    lot_column = get_first_existing_column(["LOTFROMFS", "Lot_CLASSHOT", "Lot", "LOT"], columns)
    program_column = get_first_existing_column(
        ["Program Name_CLASSHOT", "Program Name", "Program Name_RCS", "PROGRAM_NAME", "PROGRAM", "Program", "ProgramName"],
        columns,
    )

    if not visual_id_column:
        raise ValueError("Could not find a visual ID column in the input CSV.")
    if not lot_column:
        raise ValueError("Could not find a lot column in the input CSV.")

    prefix_pattern = re.compile(
        r"^ILAS_(?P<Domain>.+?)_(?P<Corner>FMIN|F\d+)_Flow(?P<Flow>[^_]+)_Freq(?P<Freq>[0-9]+(?:\.[0-9]+)?)_C(?P<CoreIndex>\d+)_MaxDTS_C$"
    )
    dts_columns: list[tuple[str, re.Match[str]]] = []
    for column in columns:
        match = prefix_pattern.match(column)
        if match:
            dts_columns.append((column, match))

    if not dts_columns:
        raise ValueError("Could not find any ILAS DTS columns ending in _MaxDTS_C.")

    lot_series = frame[lot_column].fillna("").astype(str).str.strip()
    lot_series_upper = lot_series.str.upper()

    if not program_column:
        program_column = None

    selected_program_upper = selected_program.strip().upper() if selected_program.strip() else ""
    subset = frame.copy()

    if selected_program_upper:
        if program_column is None:
            raise ValueError("The input CSV does not contain a program-name column, so it cannot be filtered by program.")
        subset = subset[subset[program_column].fillna("").astype(str).str.strip().str.upper() == selected_program_upper].copy()
        if subset.empty:
            raise ValueError(f"Program '{selected_program}' was not found in the input CSV.")
        chosen_program = selected_program.strip()
        program_lots = lot_series_upper.loc[subset.index]
        program_lots = program_lots[(program_lots != "") & (program_lots != "NAN")].unique().tolist()
        chosen_lot = ", ".join(sorted({str(item) for item in program_lots})) if program_lots else "ALL"
    else:
        chosen_program = None
        if selected_lot.strip():
            chosen_lot = selected_lot.strip()
            subset = subset[lot_series_upper.loc[subset.index] == chosen_lot.upper()].copy()
            if subset.empty:
                raise ValueError(f"Lot '{selected_lot}' was not found in the input CSV.")
        else:
            lot_counts = lot_series_upper[(lot_series_upper != "") & (lot_series_upper != "NAN")].value_counts()
            if lot_counts.empty:
                raise ValueError("The input CSV does not contain any lot values.")
            chosen_lot = str(lot_counts.index[0])
            subset = subset[lot_series_upper.loc[subset.index] == chosen_lot.upper()].copy()
        program_lots = lot_series_upper.loc[subset.index]
        program_lots = program_lots[(program_lots != "") & (program_lots != "NAN")].unique().tolist()

    if not selected_program_upper and program_column is not None:
        program_values = subset[program_column].fillna("").astype(str).str.strip()
        program_values = program_values[(program_values != "") & (program_values.str.upper() != "NAN")]
        if not program_values.empty:
            chosen_program = str(program_values.mode(dropna=True).iloc[0])

    if selected_program_upper:
        if program_lots:
            lot_names = sorted({str(item) for item in program_lots})
        else:
            lot_names = []
    else:
        lot_names = sorted({str(item) for item in program_lots})

    records: list[dict[str, object]] = []
    for max_dts_column, match in dts_columns:
        setter_column = max_dts_column.replace("_MaxDTS_C", "_Setter")
        lp_column = max_dts_column.replace("_MaxDTS_C", "_LP")
        if setter_column in subset.columns:
            setter_values = subset[setter_column]
        else:
            setter_values = pd.Series([None] * len(subset), index=subset.index)
        if lp_column in subset.columns:
            lp_values = subset[lp_column]
        else:
            lp_values = pd.Series([None] * len(subset), index=subset.index)

        dts_values = pd.to_numeric(subset[max_dts_column], errors="coerce")
        non_null = dts_values.notna()
        if not non_null.any():
            continue

        domain_token = match.group("Domain")
        corner = match.group("Corner")
        flow = match.group("Flow")
        frequency_ghz = float(match.group("Freq"))
        core_index = int(match.group("CoreIndex"))
        domain_name = resolve_domain_name(domain_token)
        frequency_label = get_display_frequency_label(frequency_ghz)

        subset_rows = subset.loc[non_null, [visual_id_column, lot_column]].copy()
        subset_rows.rename(columns={visual_id_column: "VisualID", lot_column: "LotFromFs"}, inplace=True)
        subset_rows["ProgramName"] = subset[program_column] if program_column else ""
        subset_rows["Domain"] = domain_name
        subset_rows["DomainToken"] = domain_token
        subset_rows["Corner"] = corner
        subset_rows["Flow"] = flow
        subset_rows["FrequencyGHz"] = frequency_ghz
        subset_rows["FrequencyLabel"] = frequency_label
        subset_rows["CoreIndex"] = core_index
        subset_rows["CoreLabel"] = f"core{core_index}"
        subset_rows["Setter"] = setter_values.loc[non_null].astype(str).fillna("")
        subset_rows["LimiterPattern"] = lp_values.loc[non_null].astype(str).fillna("")
        subset_rows["DtsTemperatureC"] = dts_values.loc[non_null].round(4)
        subset_rows["Module"] = [module_token_from_text(setter, limiter) for setter, limiter in zip(subset_rows["Setter"], subset_rows["LimiterPattern"])]
        records.extend(subset_rows.to_dict(orient="records"))

    if not records:
        raise ValueError("No DTS observations were found after filtering the selected lot.")

    points_df = pd.DataFrame.from_records(records)
    points_df["ChosenLot"] = chosen_lot
    points_df["LotFromFs"] = points_df["LotFromFs"].astype(str).str.strip()
    points_df["VisualID"] = points_df["VisualID"].astype(str).str.strip()
    points_df["Domain"] = points_df["Domain"].astype(str).str.strip()
    points_df["Module"] = points_df["Module"].astype(str).str.strip()

    return points_df, chosen_lot, chosen_program, lot_column, lot_names


def summarize_points(points_df: pd.DataFrame) -> pd.DataFrame:
    summaries: list[dict[str, object]] = []

    for (domain, corner, frequency_label, module), group in points_df.groupby(["Domain", "Corner", "FrequencyLabel", "Module"], sort=False):
        values = group["DtsTemperatureC"].astype(float).sort_values().to_numpy()
        if values.size == 0:
            continue

        q1 = quantile_linear(values, 0.25)
        median = quantile_linear(values, 0.5)
        q3 = quantile_linear(values, 0.75)
        iqr = q3 - q1
        lower_fence = q1 - (1.5 * iqr)
        upper_fence = q3 + (1.5 * iqr)
        in_fence = values[(values >= lower_fence) & (values <= upper_fence)]
        lower_whisker = float(in_fence[0] if in_fence.size else values[0])
        upper_whisker = float(in_fence[-1] if in_fence.size else values[-1])
        outliers = values[(values < lower_whisker) | (values > upper_whisker)]

        summaries.append(
            {
                "Domain": domain,
                "Corner": corner,
                "FrequencyLabel": frequency_label,
                "Module": module,
                "BucketLabel": f"{frequency_label} | {module}",
                "ObservationCount": int(values.size),
                "UnitCount": int(group["VisualID"].nunique()),
                "MinDtsC": round(float(values[0]), 4),
                "P25DtsC": round(q1, 4),
                "MedianDtsC": round(median, 4),
                "MeanDtsC": round(float(values.mean()), 4),
                "P75DtsC": round(q3, 4),
                "MaxDtsC": round(float(values[-1]), 4),
                "LowerWhiskerDtsC": round(lower_whisker, 4),
                "UpperWhiskerDtsC": round(upper_whisker, 4),
                "OutlierCount": int(outliers.size),
                "Above100Pct": round(float((values > 100.0).mean() * 100.0), 2),
                "CoreCoverage": "; ".join(sorted({str(core) for core in group["CoreLabel"].tolist()})),
                "Setters": "; ".join(sorted({str(setter) for setter in group["Setter"].tolist() if str(setter).strip()})),
                "LimiterPatterns": "; ".join(sorted({str(lp) for lp in group["LimiterPattern"].tolist() if str(lp).strip()})),
            }
        )

    summary_df = pd.DataFrame(summaries)
    if not summary_df.empty:
        summary_df.sort_values(["Domain", "Corner", "FrequencyLabel", "Module"], inplace=True)
    return summary_df


def plot_domain(domain_df: pd.DataFrame, domain_name: str, lot_name: str, output_path: Path) -> dict[str, object]:
    group_columns = ["Corner", "FrequencyGHz", "FrequencyLabel", "Module"]
    group_records: list[dict[str, object]] = []
    for (corner, frequency_ghz, frequency_label, module), group in domain_df.groupby(group_columns, sort=False):
        values = np.sort(group["DtsTemperatureC"].astype(float).to_numpy())
        if values.size == 0:
            continue
        group_records.append(
            {
                "corner": corner,
                "frequency_ghz": float(frequency_ghz),
                "frequency_label": frequency_label,
                "module": module,
                "label": f"{corner}\n{frequency_label} GHz\n{module}",
                "values": values,
                "unit_count": int(group["VisualID"].nunique()),
            }
        )

    if not group_records:
        raise ValueError(f"Domain {domain_name} has no plottable DTS groups.")

    group_records.sort(key=lambda item: (item["corner"].upper() != "FMIN", item["corner"], item["frequency_ghz"], module_sort_order(item["module"]), item["module"]))
    base_corner_colors = corner_color_map([str(record["corner"]) for record in group_records])

    width = max(16.0, min(30.0, 0.8 * len(group_records) + 10.0))
    height = 9.0 if len(group_records) <= 12 else 11.0
    fig, ax = plt.subplots(figsize=(width, height), constrained_layout=True)

    fig.patch.set_facecolor("white")
    ax.set_facecolor("#fbfbfa")
    ax.grid(axis="y", color="#d7dde8", linewidth=0.8, alpha=0.7)
    ax.set_axisbelow(True)

    stats = []
    positions = np.arange(1, len(group_records) + 1, dtype=float)
    jitter_rng = np.random.default_rng(42)

    for index, record in enumerate(group_records, start=1):
        values = record["values"]
        q1 = quantile_linear(values, 0.25)
        median = quantile_linear(values, 0.5)
        q3 = quantile_linear(values, 0.75)
        iqr = q3 - q1
        lower_fence = q1 - (1.5 * iqr)
        upper_fence = q3 + (1.5 * iqr)
        in_fence = values[(values >= lower_fence) & (values <= upper_fence)]
        lower_whisker = float(in_fence[0] if in_fence.size else values[0])
        upper_whisker = float(in_fence[-1] if in_fence.size else values[-1])
        outliers = values[(values < lower_whisker) | (values > upper_whisker)]

        stats.append(
            {
                "label": record["label"],
                "med": median,
                "q1": q1,
                "q3": q3,
                "whislo": lower_whisker,
                "whishi": upper_whisker,
                "fliers": outliers.tolist(),
            }
        )

        point_color = module_color_for_corner(record["corner"], record["module"], base_corner_colors)
        x_values = index + jitter_rng.normal(0.0, 0.075, size=values.size)
        ax.scatter(x_values, values, s=14, alpha=0.45, color=point_color, edgecolors="none")

    box = ax.bxp(
        stats,
        positions=positions,
        widths=0.52,
        showfliers=False,
        patch_artist=True,
        manage_ticks=False,
    )

    for index, patch in enumerate(box["boxes"]):
        patch_color = module_color_for_corner(group_records[index]["corner"], group_records[index]["module"], base_corner_colors)
        patch.set_facecolor(patch_color)
        patch.set_alpha(0.25)
        patch.set_edgecolor("#39424e")
        patch.set_linewidth(1.2)

    for whisker in box["whiskers"]:
        whisker.set_color("#39424e")
        whisker.set_linewidth(1.2)

    for cap in box["caps"]:
        cap.set_color("#39424e")
        cap.set_linewidth(1.2)

    for median_line in box["medians"]:
        median_line.set_color("#d32f2f")
        median_line.set_linewidth(1.8)

    corner_handles = [
        Patch(facecolor=module_color_for_corner(corner, "FUN", base_corner_colors), edgecolor="none", label=str(corner))
        for corner in sorted(set(str(record["corner"]) for record in group_records), key=corner_order_key)
    ]
    if corner_handles:
        legend = ax.legend(handles=corner_handles, title="Corner color rule", loc="upper right", frameon=False, fontsize=8, title_fontsize=9)
        legend._legend_box.align = "left"

    for limit_temp, label in ((90.0, "Lower limit = 90°C"), (110.0, "Upper limit = 110°C")):
        ax.axhline(limit_temp, color="red", linewidth=1.6, linestyle="--", dashes=(6, 5), alpha=0.9)
        ax.text(
            0.995,
            limit_temp,
            label,
            transform=ax.get_yaxis_transform(),
            ha="right",
            va="bottom" if limit_temp >= 100.0 else "top",
            fontsize=9,
            color="red",
            fontweight="semibold",
        )

    ax.axhline(100.0, color="#2e7d32", linewidth=2.0, linestyle="--")
    ax.text(
        0.995,
        100.0,
        "Target = 100°C",
        transform=ax.get_yaxis_transform(),
        ha="right",
        va="bottom",
        fontsize=10,
        color="#2e7d32",
        fontweight="semibold",
    )

    subtitle = f"Lot {lot_name} · Domain {domain_name} · Box+Jitter · {domain_df['VisualID'].nunique()} units · {len(domain_df):,} observations"
    ax.set_title("Test temperature dist. per corner and test", fontsize=17, fontweight="bold", pad=14)
    ax.text(0.0, 1.02, subtitle, transform=ax.transAxes, fontsize=10.5, color="#546170", ha="left", va="bottom")
    ax.set_ylabel("DTS temperature (°C)", fontsize=12)
    ax.set_xlabel("Corner → Frequency → Module", fontsize=12)
    ax.set_xticks(positions)
    ax.set_xticklabels([record["label"] for record in group_records], rotation=42, ha="right", fontsize=9)
    ax.set_xlim(0.4, len(group_records) + 0.6)

    y_min = min(record["values"].min() for record in group_records)
    y_max = max(record["values"].max() for record in group_records)
    y_span = max(5.0, y_max - y_min)
    ax.set_ylim(math.floor((y_min - 0.08 * y_span) / 5.0) * 5.0, math.ceil((max(y_max, 100.0) + 0.12 * y_span) / 5.0) * 5.0)

    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)

    fig.savefig(output_path, dpi=220, bbox_inches="tight")
    plt.close(fig)

    return {
        "domain": domain_name,
        "output_path": str(output_path),
        "group_count": len(group_records),
        "observation_count": int(domain_df.shape[0]),
    }


def build_program_ppt(output_dir: Path, image_paths: list[Path], program_name: str) -> Path:
    ppt_path = output_dir / f"dts_temperature_{safe_name(program_name)}_deck.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for image_path in image_paths:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        left = Inches(0.15)
        top = Inches(0.1)
        width = Inches(12.9)
        height = Inches(7.1)
        slide.shapes.add_picture(str(image_path), left, top, width, height)

    prs.save(str(ppt_path))
    return ppt_path


def main() -> int:
    args = parse_args()
    input_csv = Path(args.input_csv)
    if not input_csv.exists():
        raise FileNotFoundError(f"Input CSV not found: {input_csv}")

    output_dir = Path(args.output_dir) if args.output_dir.strip() else input_csv.parent / f"{input_csv.stem}_matplotlib_dts"
    output_dir.mkdir(parents=True, exist_ok=True)

    points_df, chosen_lot, chosen_program, lot_column, lot_names = build_observations(input_csv, args.lot, args.program)
    summary_df = summarize_points(points_df)

    points_path = output_dir / f"dts_temperature_points_{safe_name(chosen_lot)}.csv"
    summary_path = output_dir / f"dts_temperature_bucket_summary_{safe_name(chosen_lot)}.csv"
    manifest_path = output_dir / "run_manifest.txt"

    points_df.to_csv(points_path, index=False)
    summary_df.to_csv(summary_path, index=False)

    plot_rows: list[dict[str, object]] = []
    plot_paths: list[Path] = []
    for domain_name, domain_df in points_df.groupby("Domain", sort=False):
        domain_output = output_dir / f"dts_temperature_{safe_name(chosen_lot)}_{safe_name(domain_name)}.png"
        plot_rows.append(plot_domain(domain_df, domain_name, chosen_lot, domain_output))
        plot_paths.append(domain_output)

    plot_index_path = output_dir / f"dts_temperature_plot_index_{safe_name(chosen_lot)}.csv"
    pd.DataFrame(plot_rows).to_csv(plot_index_path, index=False)

    ppt_path = build_program_ppt(output_dir, plot_paths, chosen_program or "dts_analysis")

    with manifest_path.open("w", encoding="utf-8") as handle:
        handle.write(f"InputCsv={input_csv}\n")
        handle.write(f"LotColumn={lot_column}\n")
        handle.write(f"ChosenLot={chosen_lot}\n")
        handle.write(f"MatchedLots={'; '.join(lot_names)}\n")
        if chosen_program:
            handle.write(f"ChosenProgram={chosen_program}\n")
        handle.write(f"ObservationCount={points_df.shape[0]}\n")
        handle.write(f"DomainCount={points_df['Domain'].nunique()}\n")
        handle.write(f"SummaryCsv={summary_path}\n")
        handle.write(f"PointsCsv={points_path}\n")
        handle.write(f"PlotIndexCsv={plot_index_path}\n")
        handle.write(f"PowerPointDeck={ppt_path}\n")

    print(f"Input CSV: {input_csv}")
    print(f"Chosen lot: {chosen_lot}")
    if chosen_program:
        print(f"Chosen program: {chosen_program}")
    print(f"Output directory: {output_dir}")
    print(f"Points CSV: {points_path}")
    print(f"Summary CSV: {summary_path}")
    print(f"Plot index CSV: {plot_index_path}")
    print(f"PowerPoint deck: {ppt_path}")
    print(f"Matched lots: {', '.join(lot_names) if lot_names else 'none'}")
    for plot_row in plot_rows:
        print(f"Plot: {plot_row['output_path']}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())